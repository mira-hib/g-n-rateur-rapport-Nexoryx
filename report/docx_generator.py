"""
Générateur de rapports au format DOCX.
"""

import os
from typing import Optional
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from models.report_models import FullReport
from config.settings import settings


class DOCXReportGenerator:
    """Générateur de rapports au format DOCX avec python-docx."""

    def generate(
        self,
        report: FullReport,
        output_filename: Optional[str] = None
    ) -> str:
        """
        Génère un rapport DOCX.

        Args:
            report: Rapport complet à générer
            output_filename: Nom du fichier de sortie (optionnel)

        Returns:
            Chemin du fichier DOCX généré
        """
        # Générer le nom de fichier si non fourni
        if output_filename is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            client_name = report.cover_page.client_name.replace(" ", "_")
            output_filename = f"rapport_{client_name}_{timestamp}.docx"

        output_path = os.path.join(settings.reports_output_dir, output_filename)

        # Créer le document
        doc = Document()

        # 1. Page de garde
        self._add_cover_page(doc, report)
        doc.add_page_break()

        # 2. Résumé exécutif
        self._add_executive_summary(doc, report)
        doc.add_page_break()

        # 3. Contexte & Périmètre
        self._add_context_scope(doc, report)

        # 4. Méthodologie
        self._add_methodology(doc, report)
        doc.add_page_break()

        # 5. Analyse globale
        self._add_global_analysis(doc, report)
        doc.add_page_break()

        # 6. Findings détaillés
        self._add_findings(doc, report)
        doc.add_page_break()

        # 7. Plan d'action
        self._add_action_plan(doc, report)
        doc.add_page_break()

        # 8. Conclusion
        self._add_conclusion(doc, report)

        # Sauvegarder
        doc.save(output_path)

        return output_path

    def _add_cover_page(self, doc: Document, report: FullReport):
        """Ajoute la page de garde."""
        # Titre
        title = doc.add_heading(report.cover_page.title, level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()
        doc.add_paragraph()

        # Informations
        p = doc.add_paragraph()
        p.add_run(f"Client: ").bold = True
        p.add_run(report.cover_page.client_name)

        p = doc.add_paragraph()
        p.add_run(f"Type d'audit: ").bold = True
        p.add_run(report.cover_page.audit_type)

        p = doc.add_paragraph()
        p.add_run(f"Date: ").bold = True
        p.add_run(report.cover_page.audit_date.strftime('%d/%m/%Y'))

        p = doc.add_paragraph()
        p.add_run(f"Auditeur: ").bold = True
        p.add_run(report.cover_page.auditor)

        p = doc.add_paragraph()
        p.add_run(f"Version: ").bold = True
        p.add_run(report.cover_page.report_version)

        doc.add_paragraph()
        doc.add_paragraph()

        p = doc.add_paragraph(f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_executive_summary(self, doc: Document, report: FullReport):
        """Ajoute le résumé exécutif."""
        doc.add_heading('Résumé Exécutif', level=1)

        doc.add_paragraph(report.executive_summary.overview)

        p = doc.add_paragraph()
        p.add_run('Niveau de risque global: ').bold = True
        run = p.add_run(report.executive_summary.global_risk_level.value)
        run.font.color.rgb = self._get_severity_color_rgb(report.executive_summary.global_risk_level.value)
        run.bold = True

        doc.add_paragraph()

        # Statistiques
        table = doc.add_table(rows=5, cols=2)
        table.style = 'Light Grid Accent 1'

        table.cell(0, 0).text = 'Sévérité'
        table.cell(0, 1).text = 'Nombre'
        table.cell(1, 0).text = 'Critical'
        table.cell(1, 1).text = str(report.executive_summary.critical_count)
        table.cell(2, 0).text = 'High'
        table.cell(2, 1).text = str(report.executive_summary.high_count)
        table.cell(3, 0).text = 'Medium'
        table.cell(3, 1).text = str(report.executive_summary.medium_count)
        table.cell(4, 0).text = 'Low'
        table.cell(4, 1).text = str(report.executive_summary.low_count)

        doc.add_paragraph()

        # Recommandations
        doc.add_heading('Recommandations prioritaires:', level=3)
        for i, rec in enumerate(report.executive_summary.top_recommendations, 1):
            doc.add_paragraph(f"{i}. {rec}", style='List Number')

        doc.add_paragraph()

        p = doc.add_paragraph()
        p.add_run('Timeline de remédiation: ').bold = True
        p.add_run(report.executive_summary.remediation_timeline)

    def _add_context_scope(self, doc: Document, report: FullReport):
        """Ajoute le contexte et périmètre."""
        doc.add_heading('Contexte & Périmètre', level=1)

        doc.add_paragraph(report.context_scope.environment_description)

        p = doc.add_paragraph()
        p.add_run('Périmètre technique: ').bold = True
        p.add_run(report.context_scope.technical_scope)

        p = doc.add_paragraph()
        p.add_run('Périmètre temporel: ').bold = True
        p.add_run(report.context_scope.temporal_scope)

    def _add_methodology(self, doc: Document, report: FullReport):
        """Ajoute la méthodologie."""
        doc.add_heading('Méthodologie', level=1)

        p = doc.add_paragraph()
        p.add_run('Type d\'audit: ').bold = True
        p.add_run(report.methodology.audit_type)

        p = doc.add_paragraph()
        p.add_run('Standards utilisés: ').bold = True
        p.add_run(', '.join(report.methodology.standards))

        doc.add_paragraph(report.methodology.approach)

    def _add_global_analysis(self, doc: Document, report: FullReport):
        """Ajoute l'analyse globale."""
        doc.add_heading('Analyse Globale', level=1)

        p = doc.add_paragraph()
        p.add_run(f'Score de sécurité: ').bold = True
        p.add_run(f'{report.global_analysis.security_score}/100')

        doc.add_paragraph(report.global_analysis.score_description)

        doc.add_paragraph()

        # Distribution par sévérité
        doc.add_heading('Distribution par sévérité:', level=3)
        for severity, count in report.global_analysis.severity_distribution.items():
            doc.add_paragraph(f'• {severity}: {count}', style='List Bullet')

    def _add_findings(self, doc: Document, report: FullReport):
        """Ajoute les findings détaillés."""
        doc.add_heading('Findings Détaillés', level=1)

        for i, finding in enumerate(report.findings, 1):
            doc.add_heading(f'Finding #{i}: {finding.title}', level=2)

            p = doc.add_paragraph()
            p.add_run('Catégorie: ').bold = True
            p.add_run(finding.category.value)

            p = doc.add_paragraph()
            p.add_run('Sévérité: ').bold = True
            run = p.add_run(finding.severity.value)
            run.font.color.rgb = self._get_severity_color_rgb(finding.severity.value)
            run.bold = True

            p = doc.add_paragraph()
            p.add_run('Priorité: ').bold = True
            p.add_run(finding.priority.value)

            doc.add_paragraph()

            doc.add_heading('Description:', level=4)
            doc.add_paragraph(finding.enriched_description)

            doc.add_heading('Impact métier:', level=4)
            doc.add_paragraph(finding.business_impact)

            doc.add_heading('Recommandation:', level=4)
            doc.add_paragraph(finding.enriched_recommendation)

            if i < len(report.findings):
                doc.add_paragraph()

    def _add_action_plan(self, doc: Document, report: FullReport):
        """Ajoute le plan d'action."""
        doc.add_heading('Plan d\'Action Priorisé', level=1)

        doc.add_paragraph(report.action_plan.introduction)

        doc.add_paragraph()

        # Tableau des actions
        table = doc.add_table(rows=len(report.action_plan.actions) + 1, cols=5)
        table.style = 'Light Grid Accent 1'

        # En-tête
        table.cell(0, 0).text = '#'
        table.cell(0, 1).text = 'Mesure'
        table.cell(0, 2).text = 'Priorité'
        table.cell(0, 3).text = 'Responsable'
        table.cell(0, 4).text = 'Deadline'

        # Données
        for i, action in enumerate(report.action_plan.actions, 1):
            table.cell(i, 0).text = str(i)
            table.cell(i, 1).text = action.measure
            table.cell(i, 2).text = action.priority.value
            table.cell(i, 3).text = action.responsible
            table.cell(i, 4).text = action.deadline

    def _add_conclusion(self, doc: Document, report: FullReport):
        """Ajoute la conclusion."""
        doc.add_heading('Conclusion', level=1)

        doc.add_paragraph(report.conclusion.summary)

        doc.add_paragraph()

        doc.add_heading('Points positifs:', level=3)
        for point in report.conclusion.positive_points:
            doc.add_paragraph(f'• {point}', style='List Bullet')

        doc.add_paragraph()

        doc.add_heading('Axes d\'amélioration:', level=3)
        for area in report.conclusion.improvement_areas:
            doc.add_paragraph(f'• {area}', style='List Bullet')

        doc.add_paragraph()

        doc.add_heading('Recommandations stratégiques:', level=3)
        for rec in report.conclusion.strategic_recommendations:
            doc.add_paragraph(f'• {rec}', style='List Bullet')

        doc.add_paragraph()
        doc.add_paragraph(report.conclusion.next_steps)

    def _get_severity_color_rgb(self, severity: str) -> RGBColor:
        """Retourne la couleur RGB selon la sévérité."""
        colors_map = {
            'Critical': RGBColor(255, 0, 0),      # Rouge
            'High': RGBColor(255, 165, 0),        # Orange
            'Medium': RGBColor(255, 255, 0),      # Jaune
            'Low': RGBColor(0, 128, 0)            # Vert
        }
        return colors_map.get(severity, RGBColor(0, 0, 0))
